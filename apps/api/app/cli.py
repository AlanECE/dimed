"""CLI for DIMED API administration.

Usage:
    python -m app.cli create-admin --email admin@dimed.dz --password <pwd>
    python -m app.cli import-articles --file Articles.xlsx
"""

import argparse
import asyncio
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import uuid4

from sqlalchemy import select

from app.auth.service import hash_password
from app.db.session import async_session
from app.models.user import User, UserRole


async def create_admin(email: str, password: str) -> None:
    async with async_session() as session:
        result = await session.execute(select(User).where(User.email == email))
        if result.scalar_one_or_none():
            print(f"Admin {email} already exists, skipping.")
            return

        admin = User(
            id=uuid4(),
            email=email,
            password_hash=hash_password(password),
            role=UserRole.ADMIN,
            nom="Administrator",
            is_active=True,
        )
        session.add(admin)
        await session.commit()
        print(f"Admin {email} created successfully.")


async def import_articles(file_path: str) -> None:
    from openpyxl import load_workbook
    from sqlalchemy.dialects.postgresql import insert

    from app.models.medicament import Medicament

    path = Path(file_path)
    if not path.exists():
        print(f"Error: File not found: {file_path}")
        return

    wb = load_workbook(path, read_only=True)
    ws = wb.active

    processed = 0
    skipped = 0

    async with async_session() as db:
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if len(row) < 9:
                print(f"  Row {i}: skipped (not enough columns)")
                skipped += 1
                continue

            code_article = row[0]
            designation = row[1]
            ppa = row[6]
            fabricant = row[8]

            if not code_article or not designation:
                print(f"  Row {i}: skipped (missing code_article or designation)")
                skipped += 1
                continue

            try:
                ppa_decimal = Decimal(str(ppa)) if ppa else Decimal("0")
                if ppa_decimal <= 0:
                    raise InvalidOperation
            except (InvalidOperation, TypeError):
                print(f"  Row {i}: skipped (invalid PPA: {ppa})")
                skipped += 1
                continue

            stmt = insert(Medicament).values(
                id=uuid4(),
                code_article=str(code_article),
                designation=str(designation),
                ppa=ppa_decimal,
                fabricant=str(fabricant) if fabricant else None,
            )
            stmt = stmt.on_conflict_do_update(
                index_elements=["code_article"],
                set_={
                    "designation": stmt.excluded.designation,
                    "ppa": stmt.excluded.ppa,
                    "fabricant": stmt.excluded.fabricant,
                },
            )
            await db.execute(stmt)
            processed += 1

        await db.commit()

    wb.close()
    print(f"Import complete: {processed} processed, {skipped} skipped")


def main() -> None:
    parser = argparse.ArgumentParser(description="DIMED API CLI")
    subparsers = parser.add_subparsers(dest="command")

    create_admin_parser = subparsers.add_parser("create-admin")
    create_admin_parser.add_argument("--email", required=True)
    create_admin_parser.add_argument("--password", required=True)

    import_parser = subparsers.add_parser("import-articles")
    import_parser.add_argument("--file", required=True, help="Path to Articles.xlsx")

    args = parser.parse_args()

    if args.command == "create-admin":
        asyncio.run(create_admin(args.email, args.password))
    elif args.command == "import-articles":
        asyncio.run(import_articles(args.file))
    else:
        parser.print_help()


if __name__ == "__main__":
    main()
